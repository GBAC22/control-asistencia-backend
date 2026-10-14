from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from typing import Any
from uuid import UUID
from zipfile import BadZipFile

from django.contrib.auth.hashers import check_password, make_password
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from asistencia.data_access import repositories
from asistencia.models import (
    Asistencia,
    AuditoriaActividad,
    Estudiante,
    FechaClase,
    ModuloAcademico,
    UsuarioAdmin,
)
MAX_EXCEL_BYTES = 10 * 1024 * 1024


class ServiceError(Exception):
    def __init__(self, message: str, status_code: int = 400) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.message = message


def authenticate_student(registration_number: str, password: str = "") -> Estudiante:
    registration_number = registration_number.strip()
    if not registration_number:
        raise ServiceError("El número de registro es obligatorio.")
    password = password.strip()
    if not password:
        raise ServiceError("La contraseña es obligatoria.")
    module = repositories.get_active_module()
    if module is None:
        raise ServiceError("No hay un módulo activo.", 403)
    student = repositories.get_student_by_registration(
        registration_number, module.id_modulo
    )
    if student is None:
        raise ServiceError("Credenciales incorrectas.", 401)
    if student.password_hash:
        if not check_password(password, student.password_hash):
            raise ServiceError("Credenciales incorrectas.", 401)
    else:
        if password != student.numero_registro:
            raise ServiceError("Credenciales incorrectas.", 401)
        student.password_hash = make_password(password)
        student.save(update_fields=["password_hash"])
    return student


def authenticate_admin(email: str, password: str) -> UsuarioAdmin:
    if not email.strip() or not password:
        raise ServiceError("Correo y contraseña son obligatorios.")
    admin = repositories.get_admin_by_email(email.strip())
    if admin is None or not check_password(password, admin.password_hash):
        raise ServiceError("Credenciales incorrectas.", 401)
    return admin


def record_student_attendance(student_id: UUID) -> Asistencia:
    today = timezone.localdate()
    with transaction.atomic():
        student = repositories.get_student(student_id)
        if student is None or not student.modulo.es_activo:
            raise ServiceError("El estudiante no pertenece al módulo activo.", 403)
        module = repositories.get_module_for_update(student.modulo_id)
        if module is None or not module.es_activo:
            raise ServiceError("El estudiante no pertenece al módulo activo.", 403)
        if repositories.get_attendance_for_day(student.id_estudiante, today):
            raise ServiceError("La asistencia ya fue registrada para hoy.", 409)
        class_session = repositories.get_class_on_date(
            student.modulo_id, today, lock=True
        )
        if class_session is None or class_session.es_cancelada:
            raise ServiceError("Hoy no hay una clase programada.", 403)

        marked_at = timezone.now()
        try:
            attendance = repositories.create_attendance(
                estudiante=student,
                fecha_clase=class_session,
                estado=Asistencia.Estado.PRESENTE,
                fecha_hora_marcado=marked_at,
            )
        except IntegrityError as exc:
            raise ServiceError(
                "La asistencia ya fue registrada para esta clase.", 409
            ) from exc

        _write_audit(
            module=student.modulo,
            actor=student.numero_registro,
            role=AuditoriaActividad.RolActor.ESTUDIANTE,
            action="MARCAR_ASISTENCIA",
            entity="asistencias",
            previous=None,
            new={
                "id_asistencia": str(attendance.id_asistencia),
                "id_estudiante": str(student.id_estudiante),
                "id_fecha_clase": str(class_session.id_fecha_clase),
                "estado": attendance.estado,
                "fecha_hora_marcado": marked_at.isoformat(),
            },
        )
        return attendance


def get_today_attendance(student_id: UUID) -> dict[str, Any]:
    student = repositories.get_student(student_id)
    if student is None or not student.modulo.es_activo:
        raise ServiceError("El estudiante no pertenece al módulo activo.", 403)
    today = timezone.localdate()
    class_session = repositories.get_class_on_date(student.modulo_id, today)
    if class_session is None:
        return {"programada": False, "fecha": today.isoformat(), "asistencia": None}
    if class_session.es_cancelada:
        return {
            "programada": False,
            "cancelada": True,
            "motivo_cancelacion": class_session.motivo_cancelacion,
            "fecha": today.isoformat(),
            "hora_inicio": class_session.hora_inicio.isoformat(),
            "hora_fin": class_session.hora_fin.isoformat(),
            "asistencia": None,
        }
    attendance = repositories.get_attendance(
        student.id_estudiante, class_session.id_fecha_clase
    )
    return {
        "programada": True,
        "fecha": today.isoformat(),
        "id_fecha_clase": str(class_session.id_fecha_clase),
        "hora_inicio": class_session.hora_inicio.isoformat(),
        "hora_fin": class_session.hora_fin.isoformat(),
        "asistencia": (
            {
                "estado": attendance.estado,
                "fecha_hora_marcado": (
                    attendance.fecha_hora_marcado.isoformat()
                    if attendance.fecha_hora_marcado
                    else None
                ),
            }
            if attendance
            else None
        ),
    }


def get_student_schedule(student_id: UUID) -> list[dict[str, Any]]:
    student = repositories.get_student(student_id)
    if student is None or not student.modulo.es_activo:
        raise ServiceError("El estudiante no pertenece al módulo activo.", 403)

    today = timezone.localdate()
    schedule = []
    for class_session in repositories.get_upcoming_classes(
        student.modulo_id, today
    ):
        attendance = repositories.get_attendance(
            student.id_estudiante, class_session.id_fecha_clase
        )
        schedule.append(
            {
                "id_fecha_clase": str(class_session.id_fecha_clase),
                "fecha": class_session.fecha.isoformat(),
                "hora_inicio": class_session.hora_inicio.isoformat(),
                "hora_fin": class_session.hora_fin.isoformat(),
                "es_hoy": class_session.fecha == today,
                "es_cancelada": class_session.es_cancelada,
                "motivo_cancelacion": class_session.motivo_cancelacion,
                "asistencia": attendance.estado if attendance else None,
            }
        )
    return schedule


def update_admin_attendance(
    admin: UsuarioAdmin,
    class_id: UUID,
    student_id: UUID,
    status: str,
) -> Asistencia:
    if status not in Asistencia.Estado.values:
        raise ServiceError("El estado debe ser PRESENTE o FALTA.")

    with transaction.atomic():
        class_session = repositories.get_class_by_id(class_id)
        student = repositories.get_student(student_id)
        if class_session is None or student is None:
            raise ServiceError("No se encontró la clase o el estudiante.", 404)
        if student.modulo_id != class_session.modulo_id:
            raise ServiceError("El estudiante no pertenece al módulo de esta clase.")
        if class_session.es_cancelada:
            raise ServiceError("No se puede modificar asistencia de una clase cancelada.", 409)

        attendance = repositories.get_attendance(student_id, class_id)
        previous_status = attendance.estado if attendance else None
        previous_marked_at = attendance.fecha_hora_marcado if attendance else None
        if attendance is None:
            attendance = repositories.new_attendance(
                estudiante=student,
                fecha_clase=class_session,
            )
        attendance.estado = status
        attendance.fecha_hora_marcado = (
            attendance.fecha_hora_marcado
            if status == Asistencia.Estado.PRESENTE and attendance.fecha_hora_marcado
            else timezone.now()
            if status == Asistencia.Estado.PRESENTE
            else None
        )
        repositories.save_attendance(attendance)

        _write_audit(
            module=class_session.modulo,
            actor=admin.correo,
            role=AuditoriaActividad.RolActor.ADMIN,
            action="MODIFICAR_ASISTENCIA",
            entity="asistencias",
            previous={
                "id_estudiante": str(student_id),
                "id_fecha_clase": str(class_id),
                "estado": previous_status
                or (
                    Asistencia.Estado.FALTA
                    if class_session.fecha < timezone.localdate()
                    else "SIN_MARCA"
                ),
                "fecha_hora_marcado": (
                    previous_marked_at.isoformat() if previous_marked_at else None
                ),
            },
            new={
                "id_asistencia": str(attendance.id_asistencia),
                "id_estudiante": str(student_id),
                "id_fecha_clase": str(class_id),
                "estado": attendance.estado,
                "fecha_hora_marcado": (
                    attendance.fecha_hora_marcado.isoformat()
                    if attendance.fecha_hora_marcado
                    else None
                ),
            },
        )
        return attendance


def create_module(
    name: str, start_date: str, end_date: str
) -> ModuloAcademico:
    name = name.strip()
    if not name:
        raise ServiceError("El nombre del módulo es obligatorio.")
    try:
        start = date.fromisoformat(start_date)
        end = date.fromisoformat(end_date)
    except (TypeError, ValueError) as exc:
        raise ServiceError("Las fechas deben tener formato YYYY-MM-DD.") from exc
    if end < start:
        raise ServiceError("La fecha de fin no puede ser anterior a la de inicio.")
    return repositories.create_module(
        nombre=name,
        fecha_inicio=start,
        fecha_fin=end,
        es_activo=False,
    )


def activate_module(module_id: UUID, admin: UsuarioAdmin) -> ModuloAcademico:
    with transaction.atomic():
        modules = repositories.lock_modules_for_update()
        module = next(
            (item for item in modules if item.id_modulo == module_id), None
        )
        if module is None:
            raise ServiceError("No se encontró el módulo.", 404)
        student_count, class_count = repositories.module_import_counts(module_id)
        if student_count == 0 or class_count == 0:
            raise ServiceError(
                "Importe estudiantes y fechas de clase antes de activar el módulo.",
                409,
            )
        repositories.deactivate_other_modules(module.id_modulo)
        module.es_activo = True
        repositories.save_module_fields(module, "es_activo")
        _write_audit(
            module=module,
            actor=admin.correo,
            role=AuditoriaActividad.RolActor.ADMIN,
            action="ACTIVAR_MODULO",
            entity="modulos_academicos",
            previous=None,
            new={"id_modulo": str(module.id_modulo), "nombre": module.nombre},
        )
        return module


def import_module_workbook(
    module_id: UUID, workbook_content: bytes, admin: UsuarioAdmin
) -> dict[str, int]:
    if not workbook_content:
        raise ServiceError("El archivo Excel está vacío.")
    if len(workbook_content) > MAX_EXCEL_BYTES:
        raise ServiceError("El archivo supera el límite de 10 MB.", 413)

    try:
        workbook = load_workbook(BytesIO(workbook_content), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError) as exc:
        raise ServiceError("No se pudo leer el archivo .xlsx.") from exc

    expected_sheets = {"estudiantes", "fechas_clase"}
    if not expected_sheets.issubset(workbook.sheetnames):
        workbook.close()
        raise ServiceError(
            "El Excel debe incluir las hojas 'estudiantes' y 'fechas_clase'."
        )

    try:
        students = _read_student_sheet(workbook["estudiantes"])
        class_dates = _read_class_sheet(workbook["fechas_clase"])
    finally:
        workbook.close()

    with transaction.atomic():
        module = repositories.get_module_for_update(module_id)
        if module is None:
            raise ServiceError("No se encontró el módulo.", 404)
        if module.es_activo:
            raise ServiceError("Desactive el módulo antes de importar cambios.", 409)

        registrations = [row["numero_registro"] for row in students]
        if len(registrations) != len(set(registrations)):
            raise ServiceError("El Excel contiene números de registro duplicados.")
        existing = repositories.get_existing_registrations(
            module.id_modulo, registrations
        )
        if existing:
            raise ServiceError(
                "Ya existen estudiantes del módulo en el Excel: "
                + ", ".join(sorted(existing))
            )

        for row in class_dates:
            if not module.fecha_inicio <= row["fecha"] <= module.fecha_fin:
                raise ServiceError(
                    f"La fecha {row['fecha'].isoformat()} está fuera del rango del módulo."
                )

        repositories.bulk_create_students(module.id_modulo, students)
        repositories.bulk_create_class_dates(module.id_modulo, class_dates)
        _write_audit(
            module=module,
            actor=admin.correo,
            role=AuditoriaActividad.RolActor.ADMIN,
            action="IMPORTAR_EXCEL",
            entity="estudiantes,fechas_clase",
            previous=None,
            new={
                "estudiantes_importados": len(students),
                "fechas_clase_importadas": len(class_dates),
            },
        )
    return {"estudiantes": len(students), "fechas_clase": len(class_dates)}


def attendance_report(
    class_date: date, module_id: UUID, search: str = ""
) -> list[dict[str, Any]]:
    sessions = repositories.get_class_sessions(module_id, class_date)
    rows: list[dict[str, Any]] = []
    for class_session in sessions:
        students = repositories.get_students_for_report(module_id, search)
        existing = {
            item.estudiante_id: item
            for item in repositories.get_attendance_for_class(
                class_session.id_fecha_clase
            )
        }
        for student in students:
            attendance = existing.get(student.id_estudiante)
            rows.append(
                {
                    "id_estudiante": str(student.id_estudiante),
                    "id_fecha_clase": str(class_session.id_fecha_clase),
                    "numero_registro": student.numero_registro,
                    "nombre_completo": student.nombre_completo,
                    "correo_electronico": student.correo_electronico,
                    "fecha": class_session.fecha.isoformat(),
                    "hora_inicio": class_session.hora_inicio.isoformat(),
                    "hora_fin": class_session.hora_fin.isoformat(),
                    "estado": (
                        "CANCELADA"
                        if class_session.es_cancelada
                        else attendance.estado
                        if attendance
                        else "FALTA"
                        if class_session.fecha < timezone.localdate()
                        else "PENDIENTE"
                    ),
                    "id_asistencia": (
                        str(attendance.id_asistencia) if attendance else None
                    ),
                    "es_cancelada": class_session.es_cancelada,
                    "motivo_cancelacion": class_session.motivo_cancelacion,
                }
            )
    return rows


def set_class_cancellation(
    class_id: UUID,
    admin: UsuarioAdmin,
    *,
    cancelled: bool,
    reason: str = "",
) -> FechaClase:
    with transaction.atomic():
        existing_class = repositories.get_class_by_id(class_id)
        if existing_class is None:
            raise ServiceError("No se encontró la clase.", 404)
        module = repositories.get_module_for_update(existing_class.modulo_id)
        if module is None:
            raise ServiceError("No se encontró el módulo.", 404)
        class_session = repositories.get_class_by_id_for_update(class_id)
        if class_session is None:
            raise ServiceError("No se encontró la clase.", 404)

        if cancelled:
            reason = reason.strip()
            if not reason:
                raise ServiceError("Indica el motivo de cancelación.")
            if len(reason) > 255:
                raise ServiceError("El motivo no puede superar 255 caracteres.")
            if class_session.es_cancelada:
                raise ServiceError("La clase ya está cancelada.", 409)
            if repositories.class_attendance_count(class_id):
                raise ServiceError(
                    "No se puede cancelar una clase que ya tiene asistencias registradas.",
                    409,
                )
            action = "CANCELAR_CLASE"
            previous = {"es_cancelada": False, "motivo_cancelacion": None}
            class_session.es_cancelada = True
            class_session.motivo_cancelacion = reason
            new = {"es_cancelada": True, "motivo_cancelacion": reason}
        else:
            if not class_session.es_cancelada:
                raise ServiceError("La clase no está cancelada.", 409)
            action = "RESTAURAR_CLASE"
            previous = {
                "es_cancelada": True,
                "motivo_cancelacion": class_session.motivo_cancelacion,
            }
            class_session.es_cancelada = False
            class_session.motivo_cancelacion = None
            new = {"es_cancelada": False, "motivo_cancelacion": None}

        repositories.save_class_fields(
            class_session, "es_cancelada", "motivo_cancelacion"
        )
        _write_audit(
            module=class_session.modulo,
            actor=admin.correo,
            role=AuditoriaActividad.RolActor.ADMIN,
            action=action,
            entity="fechas_clase",
            previous={
                "id_fecha_clase": str(class_session.id_fecha_clase),
                **previous,
            },
            new={
                "id_fecha_clase": str(class_session.id_fecha_clase),
                "fecha": class_session.fecha.isoformat(),
                "hora_inicio": class_session.hora_inicio.isoformat(),
                "hora_fin": class_session.hora_fin.isoformat(),
                **new,
            },
        )
        return class_session


def _read_sheet_rows(sheet, required_headers: set[str]) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration as exc:
        raise ServiceError(f"La hoja '{sheet.title}' está vacía.") from exc
    if len(headers) != len(set(headers)):
        raise ServiceError(f"La hoja '{sheet.title}' contiene encabezados repetidos.")
    missing = required_headers - set(headers)
    if missing:
        raise ServiceError(
            f"Faltan columnas en la hoja '{sheet.title}': {', '.join(sorted(missing))}."
        )

    entries = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        entries.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
        )
    return entries


def _read_student_sheet(sheet) -> list[dict[str, Any]]:
    entries = _read_sheet_rows(sheet, {"numero_registro", "nombre_completo"})
    parsed = []
    for row_number, row in enumerate(entries, start=2):
        registration = _required_text(row["numero_registro"], "numero_registro", row_number)
        name = _required_text(row["nombre_completo"], "nombre_completo", row_number)
        email = row.get("correo_electronico")
        if email is not None and not isinstance(email, str):
            email = str(email).strip()
        pwd = row.get("contrasena") or row.get("password")
        if pwd is not None:
            pwd = str(pwd).strip()
        parsed.append(
            {
                "numero_registro": registration,
                "nombre_completo": name,
                "correo_electronico": email.strip() if email else None,
                "password_hash": make_password(pwd if pwd else registration),
            }
        )
    if not parsed:
        raise ServiceError("La hoja 'estudiantes' no contiene estudiantes.")
    return parsed


def _read_class_sheet(sheet) -> list[dict[str, Any]]:
    entries = _read_sheet_rows(sheet, {"fecha", "hora_inicio", "hora_fin"})
    parsed = []
    for row_number, row in enumerate(entries, start=2):
        class_date = _parse_date(row["fecha"], row_number)
        start = _parse_time(row["hora_inicio"], "hora_inicio", row_number)
        end = _parse_time(row["hora_fin"], "hora_fin", row_number)
        if end <= start:
            raise ServiceError(f"La hora_fin debe ser posterior en la fila {row_number}.")
        parsed.append({"fecha": class_date, "hora_inicio": start, "hora_fin": end})
    if not parsed:
        raise ServiceError("La hoja 'fechas_clase' no contiene fechas.")
    return parsed


def _required_text(value: Any, field: str, row_number: int) -> str:
    if value is None or not str(value).strip():
        raise ServiceError(f"Falta '{field}' en la fila {row_number}.")
    return str(value).strip()


def _parse_date(value: Any, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ServiceError(f"Fecha inválida en fila {row_number}; use YYYY-MM-DD.") from exc


def _parse_time(value: Any, field: str, row_number: int) -> time:
    if isinstance(value, time):
        return value.replace(tzinfo=None)
    try:
        return time.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ServiceError(
            f"Hora inválida en '{field}', fila {row_number}; use HH:MM."
        ) from exc


def _write_audit(
    *,
    module: ModuloAcademico,
    actor: str,
    role: str,
    action: str,
    entity: str,
    previous: dict[str, Any] | None,
    new: dict[str, Any] | None,
) -> None:
    repositories.create_audit(
        modulo=module,
        usuario_actor=actor,
        rol_actor=role,
        accion=action,
        entidad_afectada=entity,
        valor_anterior=previous,
        valor_nuevo=new,
    )


def change_student_password(
    student_id: UUID, current_password: str, new_password: str
) -> None:
    current_password = current_password.strip()
    new_password = new_password.strip()
    if not current_password or not new_password:
        raise ServiceError("La contraseña actual y la nueva son obligatorias.")
    if len(new_password) < 4:
        raise ServiceError("La nueva contraseña debe tener al menos 4 caracteres.")

    with transaction.atomic():
        student = repositories.get_student(student_id)
        if student is None or not student.modulo.es_activo:
            raise ServiceError("El estudiante no pertenece al módulo activo.", 403)
        if student.password_hash:
            if not check_password(current_password, student.password_hash):
                raise ServiceError("La contraseña actual es incorrecta.", 401)
        else:
            if current_password != student.numero_registro:
                raise ServiceError("La contraseña actual es incorrecta.", 401)

        student.password_hash = make_password(new_password)
        student.save(update_fields=["password_hash"])
        _write_audit(
            module=student.modulo,
            actor=student.numero_registro,
            role=AuditoriaActividad.RolActor.ESTUDIANTE,
            action="CAMBIAR_PASSWORD",
            entity="estudiantes",
            previous=None,
            new={"id_estudiante": str(student.id_estudiante), "numero_registro": student.numero_registro},
        )
